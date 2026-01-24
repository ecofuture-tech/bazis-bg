

import os
from datetime import date, datetime
from io import BytesIO
from tempfile import NamedTemporaryFile
from uuid import UUID
from zipfile import ZIP_DEFLATED, ZipFile

from django.db.models import Manager, Model
from django.utils.functional import Promise
from django.utils.timezone import now

from openpyxl import Workbook, styles
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from ..basic.base import BgBase


WRITABLE_HEAD_FONT = styles.Font(bold=True)
READABLE_FONT = styles.Font(sz=9)


class WBook:
    row_num = 0
    time = None
    _counter = None

    @classmethod
    def reset(cls):
        cls._counter = 0
        cls.time = now().strftime('%Y-%m-%d_%H-%M-%S')

    def __init__(self, parent):
        WBook._counter += 1

        # file object
        self.fp = BytesIO()
        self.parent = parent
        self.counter = WBook._counter
        self.unions = set()
        self.fields_read = parent.get_fields_read()
        self.fields_marked = parent.get_fields_marked()

        if parent.write_only:
            self.wb = Workbook(write_only=True)
            self.ws = self.wb.create_sheet()
        else:
            self.wb = Workbook()
            self.ws = self.wb.active

        titles = parent.get_titles()
        groups = []

        # prepare the header matrix
        gen_row_title = lambda: [None] * len(self.fields_read)  # NOQA [E731]
        if 'groups' in titles and isinstance(titles['groups'], list):
            rows_title = [gen_row_title(), gen_row_title()]
            groups = titles['groups']
            titles = titles['titles']
        else:
            rows_title = [gen_row_title()]

        for i, f_name in enumerate(self.fields_read):
            letter = get_column_letter(i + 1)
            # if column sizes are set - use them
            self.ws.column_dimensions[letter].width = parent.fields_size.get(f_name, 20)
            self.ws.column_dimensions[letter].auto_size = True
            # write Russian headers, if they exist
            if titles:
                # field cell
                cell = self.title_cell_bild(
                    titles.get(f_name, f_name), f_name in self.fields_marked
                )

                if groups:
                    # find the group for the current column
                    group = ([group for group in groups if f_name in group['columns']] + [None])[0]
                    if group:
                        rows_title[0][i] = self.title_cell_bild(
                            group['name_union'], group.get('marked')
                        )
                        rows_title[1][i] = cell
                        self.unions.add(
                            '{}1:{}1'.format(
                                self.get_letter_for_field(group['columns'][0]),
                                self.get_letter_for_field(group['columns'][-1]),
                            )
                        )
                    else:
                        rows_title[0][i] = cell
                        self.unions.add(f'{letter}1:{letter}2')
                else:
                    rows_title[0][i] = cell

        for row in rows_title:
            self.append(row)

    def title_cell_bild(self, value, marked=False):
        # convert lazy object to string in value
        c = WriteOnlyCell(ws=self.ws, value=str(value))
        if marked:
            c.font = WRITABLE_HEAD_FONT
        else:
            c.font = READABLE_FONT
        c.alignment = Alignment(wrapText=True, horizontal='center', vertical='center')
        return c

    def get_letter_for_field(self, field):
        try:
            return get_column_letter(self.fields_read.index(field) + 1)
        except ValueError:
            return None

    def append(self, item):
        self.row_num += 1
        self.ws.append(item)

    def get_name(self):
        return f'{self.parent.file_name}.{self.time}.xlsx'

    def save(self):
        for union in self.unions:
            self.ws.merged_cells.ranges.append(union)
        self.parent.pre_save()
        self.wb.save(self.fp)


class Zip:

    def __init__(self):
        self.zp = None
        self.fp = None

    def write(self, wbook):
        if not self.zp:
            self.fp = NamedTemporaryFile(mode='w+', prefix='bg_download_zip.', delete=False)
            self.zp = ZipFile(self.fp, mode="w", compression=ZIP_DEFLATED, allowZip64=True)

        self.zp.writestr(f'{wbook.counter}.{wbook.get_name()}', wbook.fp.getvalue())

    @property
    def is_init(self):
        return bool(self.zp)

    def close(self):
        self.zp.close()

    def unlink(self):
        if os.path.exists(self.fp.name):
            os.remove(self.fp.name)


class BgBaseDownload(BgBase):
    write_only = True
    is_data_ext = False
    fields_read = []
    fields_marked = set()
    fields_size = {}
    file_name = ''
    limit_rows_split = 500000
    parallel = 100

    def get_titles(self):
        return {}

    def get_fields_read(self):
        return self.fields_read

    def get_fields_marked(self):
        return self.fields_marked

    def get_queryset(self):
        return []

    def get_count(self):
        raise Exception('A specific number of results is expected')

    def pre_save(self):
        pass

    def handle(self):
        count = self.get_count()

        # presets
        WBook.reset()

        self.next_phase('Export', count)

        # references to structures
        fields_read = self.get_fields_read()
        qs = self.get_queryset()

        # create a new workbook
        self.wbook = WBook(self)

        # archive object
        _zip = Zip()

        # total number of rows
        total_count = 0
        # write each row
        row_count = 0

        for _i, row in enumerate(qs):
            if row_count >= self.limit_rows_split:
                # save data
                self.wbook.save()
                # write to zip
                _zip.write(self.wbook)
                # counters
                row_count = 0
                # create a new workbook
                self.wbook = WBook(self)

            # get a dictionary of values from the queryset record
            items = {
                k: (row.get(k, None) if isinstance(row, dict) else getattr(row, k, None))
                for k in fields_read
            }
            # standard values
            data = {}
            # extended values (can only exist if the corresponding key is the last one)
            data_ext = []

            # normalize
            for k, v in items.items():
                if isinstance(v, (str, Model, UUID)):
                    data[k] = str(v)
                elif isinstance(v, bool):
                    data[k] = int(v)
                elif isinstance(v, (Manager, list)):
                    # field values
                    if isinstance(v, Manager):
                        v_nested = [str(v_ext) for v_ext in v.all()]
                    else:
                        v_nested = v

                    if self.is_data_ext and k == fields_read[-1]:
                        data[k] = ''
                        data_ext = v_nested
                    else:
                        data[k] = ', '.join([str(x) for x in v_nested])
                elif isinstance(v, datetime):
                    data[k] = v.strftime('%Y-%m-%d %H:%M:%S')
                elif isinstance(v, date):
                    data[k] = v.strftime('%d.%m.%Y')
                elif isinstance(v, memoryview):
                    data[k] = str(v.tobytes())
                else:
                    data[k] = str(v)

            # assemble the list of values
            data = [data[key] for key in fields_read]
            if data_ext:
                data = data[:-1] + data_ext

            self.wbook.append(data)
            row_count += 1
            total_count += 1

            self.progress(d_count=1000)

        self.wbook.save()

        # if the archive is initialized
        if _zip.is_init:
            # write to the archive
            _zip.write(self.wbook)
            _zip.close()
            # and write it to the task as well
            self.task_file_save(
                f'{self.file_name}.{WBook.time}.zip',
                _zip.fp,
                file_params={
                    'fields': fields_read,
                    'titles': {
                        x: str(y) if isinstance(y, Promise) else y
                        for x, y in self.get_titles().items()
                    },
                    'total_count': total_count,
                },
            )
            # delete the temporary archive
            _zip.unlink()
        else:
            self.task_file_save(
                self.wbook.get_name(),
                self.wbook.fp,
                file_params={
                    'fields': fields_read,
                    'titles': {
                        x: str(y) if isinstance(y, Promise) else y
                        for x, y in self.get_titles().items()
                    },
                    'total_count': total_count,
                },
            )
